"""Печать готовых таблиц и диаграмм в CSV, XLSX и PDF без запросов к БД."""
import csv
import io
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
from fastapi import HTTPException
from openpyxl import Workbook
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer,KeepTogether
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4,landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.graphics.shapes import Drawing,Rect,String,Line,PolyLine,Circle
from .resources import LABELS
from .report_layout import REPORT_COLUMNS,REPORT_LABELS,REPORT_TITLES


def safe_cell(value):
    if value is None: return ''
    if isinstance(value,str) and value[:1] in ('=','+','-','@','\t','\r'): return "'"+value
    return value

def render_export(kind,fmt,rows,context=None):
    columns=REPORT_COLUMNS[kind]
    headers=[REPORT_LABELS.get(c,LABELS.get(c,{'EquipmentName':'Оборудование','WorkshopName':'Цех'}.get(c,c))) for c in columns]
    if fmt=='csv':
        out=io.StringIO();writer=csv.writer(out,delimiter=';');writer.writerow(headers)
        writer.writerows([[safe_cell(r.get(c)) for c in columns] for r in rows])
        return out.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    if fmt=='xlsx':
        wb=Workbook();ws=wb.active;ws.title='Отчёт';ws.append(headers)
        for row in rows: ws.append([safe_cell(row.get(c)) for c in columns])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        from openpyxl.styles import Font,PatternFill
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='153047')
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(45,max(16,max(len(str(c.value or '')) for c in col)+2))
        output=io.BytesIO();wb.save(output);return output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    if fmt=='pdf':
        font=Path('C:/Windows/Fonts/arial.ttf')
        if not font.exists(): font=Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
        if 'ToirFont' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('ToirFont',str(font)))
        styles=getSampleStyleSheet()
        for style in styles.byName.values(): style.fontName='ToirFont'
        styles['BodyText'].fontSize=8;styles['BodyText'].leading=11
        output=io.BytesIO();doc=SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=28,bottomMargin=28)
        matrix=[[Paragraph(escape(str(h)),styles['BodyText']) for h in headers]]
        matrix += [[Paragraph(escape(str(r.get(c) if r.get(c) is not None else '')),styles['BodyText']) for c in columns] for r in rows]
        grid=Table(matrix,repeatRows=1,colWidths=[(landscape(A4)[0]-48)/len(columns)]*len(columns))
        grid.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dbeafe')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f1f5f9')]),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        elements=[Paragraph('ТОиР-Проф · '+REPORT_TITLES[kind],styles['Title']),
                  Paragraph('Сформирован '+datetime.utcnow().strftime('%d.%m.%Y %H:%M UTC'),styles['Normal'])]
        context=context or {}
        report_period=context.get('summary',{}).get('period')
        if report_period:
            elements.append(Paragraph(escape(f"Период UTC: {report_period['from']} <= время < {report_period['to']}"),styles['BodyText']))
        filter_labels={'from':'С даты','to':'По дату','workshop_id':'Цех ID','workshop_ids':'Цеха ID',
            'equipment_id':'Оборудование ID','order_type':'Тип заявки','priority':'Приоритет','block':'Блок'}
        filters_text='; '.join(f'{filter_labels.get(name,name)}: {value}' for name,value in context.get('filters',{}).items() if value not in (None,''))
        elements.append(Paragraph(escape('Фильтры: '+(filters_text or 'все доступные записи по правам пользователя')),styles['BodyText']))
        elements.extend([Spacer(1,14),grid])
        for chart in context.get('charts',[]):
            items=chart['items']
            if not items: continue
            heading=[Spacer(1,16),Paragraph(escape(chart['title']+' ('+chart['unit']+')'),styles['Heading2'])]
            if chart.get('mode')=='line':
                drawing=Drawing(770,200)
                maximum=max(float(row['value'] or 0) for row in items) or 1
                points=[]
                for index,row in enumerate(items):
                    x=35+index*700/max(1,len(items)-1);y=30+float(row['value'] or 0)/maximum*145
                    points.extend([x,y])
                    drawing.add(Circle(x,y,2.5,fillColor=colors.HexColor('#2b7b9c'),strokeColor=None))
                    if index==0 or index==len(items)-1 or index%5==0:
                        drawing.add(String(x,12,str(row['label'])[5:],fontName='ToirFont',fontSize=8,textAnchor='middle'))
                if len(items)>1: drawing.add(PolyLine(points,strokeColor=colors.HexColor('#2b7b9c'),strokeWidth=1.5))
                drawing.add(Line(35,30,735,30,strokeColor=colors.HexColor('#cbd5e1')))
                drawing.add(String(25,30,'0',fontName='ToirFont',fontSize=8,textAnchor='end'))
                drawing.add(String(25,175,f'{maximum:g}',fontName='ToirFont',fontSize=8,textAnchor='end'))
                elements.append(KeepTogether(heading+[drawing]))
                continue
            # Labeled horizontal bars avoid rotated Cyrillic labels and preserve zero values.
            drawing=Drawing(770,len(items)*21+20)
            maximum=max(float(row['value'] or 0) for row in items) or 1
            for index,row in enumerate(items):
                y=(len(items)-index-1)*21+8
                label=str(row['label'] or '')
                if len(label)>34: label=label[:31]+'...'
                value=float(row['value'] or 0)
                drawing.add(String(0,y+3,label,fontName='ToirFont',fontSize=8))
                drawing.add(Rect(210,y,max(0,value)/maximum*485,13,fillColor=colors.HexColor('#2b7b9c'),strokeColor=None))
                drawing.add(String(710,y+3,f'{value:g}',fontName='ToirFont',fontSize=8))
            elements.append(KeepTogether(heading+[drawing]))
        def footer(canvas,document):
            canvas.setFont('ToirFont',8);canvas.setFillColor(colors.HexColor('#64748b'))
            canvas.drawRightString(landscape(A4)[0]-24,14,f'Страница {document.page}')
        doc.build(elements,onFirstPage=footer,onLaterPages=footer)
        return output.getvalue(),'application/pdf'
    raise HTTPException(400,'Формат должен быть pdf, xlsx или csv.')
