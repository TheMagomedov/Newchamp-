"""Own random fixtures only; never resets or edits the supplied training dataset."""
import asyncio
import io
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime,timedelta
import pytest
from sqlalchemy import select
from openpyxl import load_workbook
from fastapi import HTTPException
from fastapi.testclient import TestClient
from app.main import app,limits
from app.config import DATABASE
from app.db import table,transaction
from app.advanced import revision_state,event_stream


@pytest.fixture(scope='module')
def client():
    assert DATABASE.startswith('ChampTraining_')
    with TestClient(app,raise_server_exceptions=False) as value: yield value


@pytest.fixture(autouse=True)
def reset_limit():
    limits.clear()


def login(client,role='admin'):
    response=client.post('/api/auth/login',json={'email':role+'@toir.local','password':'ChampDemo2027!'})
    assert response.status_code==200,response.text
    data=response.json()['data']
    return {'Authorization':'Bearer '+data['access_token']},data['user']


@pytest.fixture
def own(client):
    headers,user=login(client,'master')
    prefix='ADV-'+uuid.uuid4().hex[:12]
    ids={};start=datetime.utcnow().replace(microsecond=0)-timedelta(days=2)
    try:
        with transaction(user['Id']) as conn:
            def add(name,**values):
                t=table(name);ident=conn.execute(t.insert().values(**values)).inserted_primary_key[0]
                ids.setdefault(name,[]).append(ident);return ident
            team=add('Teams',Code=prefix,Name=prefix,Specialization='Test',WorkshopId=user['WorkshopId'])
            equipment=add('Equipment',InventoryNumber=prefix,Name=prefix,Model='Precision model',
                WorkshopId=user['WorkshopId'],CommissionedOn='2020-01-01',StartedOn='2020-01-02')
            material=add('Materials',Code=prefix,Name=prefix,Unit='кг',Stock=10,MinimumStock=2,UnitPrice=1.20,Location='Regression shelf')
            order=add('WorkOrders',EquipmentId=equipment,MasterId=user['Id'],TeamId=team,Type='текущая',Priority='высокий',
                Description=prefix,CreatedAt=start,PlannedStart=start+timedelta(hours=1),PlannedEnd=start+timedelta(hours=3))
            orders=table('WorkOrders')
            conn.execute(orders.update().where(orders.c.Id==order).values(Status='назначена'))
            conn.execute(orders.update().where(orders.c.Id==order).values(Status='в работе',ActualStart=start+timedelta(hours=1)))
            mtype=conn.execute(select(table('MaintenanceTypes').c.Id).where(table('MaintenanceTypes').c.IsActive==True)).first()[0]
            planned=datetime.utcnow().replace(microsecond=0)+timedelta(days=2)
            schedule=add('MaintenanceSchedule',EquipmentId=equipment,MaintenanceTypeId=mtype,TeamId=team,PlannedAt=planned)
        yield dict(headers=headers,user=user,prefix=prefix,ids=ids,equipment=equipment,team=team,material=material,
                   order=order,schedule=schedule,start=start,planned=planned)
    finally:
        with transaction() as conn:
            # Remove generated order of only this fixture's schedule before its parent.
            schedules=ids.get('MaintenanceSchedule',[])
            if schedules:
                orders=table('WorkOrders');conn.execute(orders.delete().where(orders.c.ScheduleId.in_(schedules)))
            for name in ('WorkOrders','MaintenanceSchedule','Equipment','Teams','Materials'):
                if ids.get(name):
                    t=table(name);conn.execute(t.delete().where(t.c.Id.in_(ids[name])))


def test_rich_details_filters_search_notifications_are_scoped(client,own):
    headers=own['headers'];eq=own['equipment'];order=own['order']
    result=client.get('/api/analytics/dashboard',headers=headers,params={'equipment_id':eq,'from':own['start'].date().isoformat()})
    assert result.status_code==200,result.text
    data=result.json()['data']
    assert data['metrics']['equipment']==1 and data['metrics']['open_orders']==1
    assert data['metrics']['downtime_hours']>40
    assert data['orders_by_priority']==[{'label':'высокий','value':1}]
    wrong=client.get('/api/work-orders',headers=headers,params={'equipment_id':eq,'priority':'низкий'})
    assert wrong.json()['data']['total']==0
    correct=client.get('/api/work-orders',headers=headers,params={'equipment_id':eq,'priority':'высокий','type':'текущая','date_from':own['start'].date().isoformat()})
    assert [row['Id'] for row in correct.json()['data']['items']]==[order]
    assert client.get('/api/equipment',headers=headers,params={'search':own['prefix'],'model':'absent'}).json()['data']['total']==0
    assert client.get('/api/materials',headers=headers,params={'search':own['prefix'],'location':'Regression'}).json()['data']['total']==1
    detail=client.get(f'/api/equipment/{eq}/details',headers=headers).json()['data']
    assert [row['Id'] for row in detail['orders']]==[order] and detail['history']
    assert [row['Id'] for row in detail['schedule']]==[own['schedule']]
    team=client.get(f'/api/teams/{own["team"]}',headers=headers).json()['data']
    assert team['workload']['active']==1 and team['workload']['planned_hours']==2
    assert team['members']==[]
    notifications=client.get('/api/notifications',headers=headers).json()['data']['items']
    assert any(row['id']==f'order-overdue:{order}' for row in notifications)
    assert client.get('/api/search',headers=headers,params={'q':own['prefix']}).json()['data']['total']>=3
    member,_=login(client,'member')
    assert client.get('/api/search',headers=member,params={'q':own['prefix']}).json()['data']['items']==[
        {'resource':'materials','id':own['material'],'label':own['prefix'],'subtitle':own['prefix']}]
    assert client.get(f'/api/work-orders/{order}',headers=member).status_code==404
    assert not any(row['entity_id']==order and row['resource']=='work-orders' for row in client.get('/api/notifications',headers=member).json()['data']['items'])
    scoped=client.get('/api/analytics/dashboard',headers=member,params={'equipment_id':eq}).json()['data']
    assert scoped['metrics']['equipment']==0 and scoped['metrics']['open_orders']==0


def test_schedule_generation_is_concurrency_safe_across_different_keys(client,own):
    endpoint=f'/api/schedule/{own["schedule"]}/create-order'
    payload={'master_id':own['user']['Id'],'planned_hours':2}
    keys=[str(uuid.uuid4()),str(uuid.uuid4())]
    def send(key): return client.post(endpoint,headers={**own['headers'],'Idempotency-Key':key},json=payload)
    with ThreadPoolExecutor(max_workers=2) as executor: responses=list(executor.map(send,keys))
    assert all(response.status_code==201 for response in responses),[response.text for response in responses]
    orders=[response.json()['data'] for response in responses]
    assert orders[0]['Id']==orders[1]['Id']
    assert sorted(row['ScheduleOrderReused'] for row in orders)==[False,True]
    assert orders[0]['Status']=='назначена'
    assert orders[0]['DisplayNumber'].startswith('Z-')
    assert send(keys[0]).json()==responses[0].json()
    filtered=client.get('/api/schedule',headers=own['headers'],params={'equipment_id':own['equipment'],
        'date_from':own['planned'].date().isoformat(),'date_to':own['planned'].date().isoformat()})
    assert filtered.json()['data']['total']==1
    upcoming=client.get('/api/schedule/upcoming',headers=own['headers'],params={'equipment_id':own['equipment'],'days':3})
    assert [row['Id'] for row in upcoming.json()['data']['items']]==[own['schedule']]
    manager,_=login(client,'manager')
    assert client.post(endpoint,headers=manager,json=payload).status_code==403


def test_atomic_completion_rolls_back_materials_on_insufficient_stock(client,own):
    endpoint=f'/api/work-orders/{own["order"]}/complete'
    payload={'completion_report':'Проверено, исправно','actual_end':(datetime.utcnow()-timedelta(minutes=1)).isoformat(),
        'materials':[{'MaterialId':own['material'],'Quantity':11}]}
    rejected=client.post(endpoint,headers=own['headers'],json=payload)
    assert rejected.status_code==409,rejected.text
    detail=client.get(f'/api/work-orders/{own["order"]}',headers=own['headers']).json()['data']
    assert detail['Status']=='в работе' and detail['materials']==[]
    assert client.get(f'/api/materials/{own["material"]}',headers=own['headers']).json()['data']['Stock']==10
    engineer,_=login(client,'engineer')
    assert client.post(endpoint,headers=engineer,json=payload).status_code==403
    payload['materials'][0]['Quantity']=0.125
    headers={**own['headers'],'Idempotency-Key':str(uuid.uuid4())}
    response=client.post(endpoint,headers=headers,json=payload)
    assert response.status_code==200,response.text
    assert client.post(endpoint,headers=headers,json=payload).json()==response.json()
    detail=client.get(f'/api/work-orders/{own["order"]}',headers=own['headers']).json()['data']
    assert detail['Status']=='выполнена' and len(detail['materials'])==1
    assert detail['materials'][0]['UnitPrice']==1.2
    assert detail['materials'][0]['Quantity']==0.125
    assert client.get(f'/api/materials/{own["material"]}',headers=own['headers']).json()['data']['Stock']==9.875


@pytest.mark.parametrize('kind,fmt,signature',[
    ('downtime','csv',b'\xef\xbb\xbf'),('schedule','xlsx',b'PK'),('team-efficiency','pdf',b'%PDF'),
    ('material-usage','csv',b'\xef\xbb\xbf'),('analytics','pdf',b'%PDF'),('analytics','xlsx',b'PK')])
def test_analytic_previews_exports_match_contract(client,own,kind,fmt,signature,tmp_path):
    params={'type':kind,'equipment_id':own['equipment'],'from':own['start'].date().isoformat()}
    preview=client.get('/api/reports/preview',headers=own['headers'],params=params)
    assert preview.status_code==200,preview.text
    data=preview.json()['data'];assert data['columns'] and isinstance(data['rows'],list)
    response=client.post('/api/reports/export',headers=own['headers'],json={**params,'format':fmt})
    assert response.status_code==200,response.text
    downloaded=client.get(response.json()['data']['url'],headers=own['headers'])
    assert downloaded.status_code==200 and downloaded.content.startswith(signature)
    if fmt=='pdf':
        artifact=tmp_path/f'{kind}-preview.pdf'
        artifact.write_bytes(downloaded.content)
        print('PDF visual QA:',artifact)
    if fmt=='xlsx':
        workbook=load_workbook(io.BytesIO(downloaded.content));assert workbook.active.max_row==len(data['rows'])+1
        workbook.close()
    if kind=='downtime': assert [row['inventory_number'] for row in data['rows']]==[own['prefix']]


def test_notification_revision_changes_only_for_visible_data(client,own):
    member,user=login(client,'member')
    before_master,_=revision_state(own['user']);before_member,_=revision_state(user)
    response=client.put(f'/api/equipment/{own["equipment"]}',headers=own['headers'],json={'Name':own['prefix']+' changed'})
    assert response.status_code==200,response.text
    after_master,_=revision_state(own['user']);after_member,_=revision_state(user)
    assert before_master!=after_master
    assert before_member==after_member
    assert client.get('/api/events').status_code==401


def test_sse_first_event_and_revoked_session_end_stream(monkeypatch):
    import app.advanced as module
    monkeypatch.setattr(module,'revision_state',lambda user:('abc',{'equipment':[1,1,1]}))
    def expired(credentials): raise HTTPException(401,'expired')
    monkeypatch.setattr(module,'current_user',expired)
    class Connected:
        async def is_disconnected(self): return False
    async def run():
        stream=event_stream(Connected(),{'Id':1},object(),poll_seconds=0)
        first=await anext(stream);second=await anext(stream)
        assert 'event: revision' in first and 'equipment' in first
        assert 'event: auth_expired' in second
        with pytest.raises(StopAsyncIteration): await anext(stream)
    asyncio.run(run())


def test_master_lookup_and_search_by_name_and_display_number(client,own):
    headers=own['headers'];admin,admin_user=login(client)
    lookups=client.get('/api/lookups',headers=headers).json()['data']
    assert lookups['masters']==[{'id':own['user']['Id'],'label':' '.join(own['user']['Name'].split()[:2]),
                                'workshop_id':own['user']['WorkshopId']}]
    order=client.get(f'/api/work-orders/{own["order"]}',headers=headers).json()['data']
    for phrase in (own['user']['Name'],order['DisplayNumber'],f'{order["SequenceNumber"]:04d}'):
        results=client.get('/api/search',headers=headers,params={'q':phrase,'limit':100})
        assert results.status_code==200,results.text
        assert any(row['resource']=='work-orders' and row['id']==own['order'] for row in results.json()['data']['items'])
    invalid=client.put(f'/api/work-orders/{own["order"]}',headers=admin,json={'MasterId':admin_user['Id']})
    assert invalid.status_code==400,invalid.text


def test_expired_export_returns_gone_but_is_not_visible_to_another_user(client,monkeypatch):
    import app.reports as reports
    headers,_=login(client)
    result=client.post('/api/reports/export',headers=headers,json={'type':'equipment','format':'csv'})
    assert result.status_code==200,result.text
    url=result.json()['data']['url'];token=url.rsplit('/',1)[1]
    other,_=login(client,'master')
    with reports.download_lock: reports.downloads[token]['expires']=0
    assert client.get(url,headers=other).status_code==404
    assert client.get(url,headers=headers).status_code==410
